"""Pruebas unitarias y de integracion para CU31: Generar reportes ejecutivos y consultas por voz."""

from datetime import datetime, timezone
import io
from typing import Generator
from unittest.mock import MagicMock

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from core.database import get_db
from core.security import create_access_token
from main import app
from modules.autenticacion_seguridad.modelos import UsuarioORM
from modules.comercial.cu31_reportes_voz.errores import (
    AccesoReporteDenegadoError,
    ComandoVozNoReconocidoError,
)
from modules.comercial.cu31_reportes_voz.esquemas import (
    ComandoVozIn,
    FormatoReporteEnum,
    IntencionVozEnum,
    ModuloReporteEnum,
    RangoTemporalEnum,
    ReporteFiltrosIn,
)
from modules.comercial.cu31_reportes_voz.generadores.generador_csv import GeneradorCSV
from modules.comercial.cu31_reportes_voz.generadores.generador_excel import GeneradorExcel
from modules.comercial.cu31_reportes_voz.generadores.generador_pdf import GeneradorPDF
from modules.comercial.cu31_reportes_voz.parser_voz import ParserComandosVoz
from modules.comercial.cu31_reportes_voz.servicio import ServicioReportesVoz
from modules.gestion_operativa.modelos import SucursalORM


# --- Fixtures Auxiliares ---


@pytest.fixture
def usuario_admin() -> UsuarioORM:
    return UsuarioORM(
        id_usuario=1,
        email="admin@fashionstore.com",
        password_hash="fakehash",
        nombres="Super",
        apellidos="Admin",
        rol="administrador",
        activo=True,
    )


@pytest.fixture
def usuario_encargado() -> UsuarioORM:
    return UsuarioORM(
        id_usuario=2,
        email="encargado@fashionstore.com",
        password_hash="fakehash",
        nombres="Carlos",
        apellidos="Mendoza",
        rol="encargado_sucursal",
        id_sucursal=10,
        activo=True,
    )


@pytest.fixture
def usuario_cajero() -> UsuarioORM:
    return UsuarioORM(
        id_usuario=3,
        email="cajero@fashionstore.com",
        password_hash="fakehash",
        nombres="Ana",
        apellidos="Rios",
        rol="cajero",
        id_sucursal=10,
        activo=True,
    )


@pytest.fixture
def token_admin(usuario_admin: UsuarioORM) -> str:
    return create_access_token({"sub": str(usuario_admin.id_usuario), "rol": usuario_admin.rol})


@pytest.fixture
def token_encargado(usuario_encargado: UsuarioORM) -> str:
    return create_access_token({"sub": str(usuario_encargado.id_usuario), "rol": usuario_encargado.rol})


@pytest.fixture
def token_cajero(usuario_cajero: UsuarioORM) -> str:
    return create_access_token({"sub": str(usuario_cajero.id_usuario), "rol": usuario_cajero.rol})


@pytest.fixture
def mock_db() -> MagicMock:
    db = MagicMock(spec=Session)
    # Simular scalar y scalars
    db.scalar.return_value = 5
    first_mock = MagicMock()
    first_mock.__getitem__.side_effect = lambda idx: 10 if idx == 0 else 1500.0
    db.execute.return_value.first.return_value = first_mock
    db.execute.return_value.scalars.return_value.all.return_value = []
    db.execute.return_value.all.return_value = [(10, "Equipetrol"), (20, "Calacoto")]
    db.get.return_value = SucursalORM(id_sucursal=10, nombre="Equipetrol", activa=True)
    return db


# =========================================================================
# 1. Pruebas Unitarias de Generadores Binarios
# =========================================================================


def test_generador_excel_retorna_bytes_validos():
    """Valida que GeneradorExcel construya un libro XLSX legible con openpyxl."""
    cols = ["CODIGO", "CLIENTE", "MONTO"]
    filas = [
        ["V-001", "Ana Valenzuela", 1250.50],
        ["V-002", "Carlos Mendez", 850.00],
    ]
    totales = {"MONTO": 2100.50}
    meta = {"Sucursal": "Central"}

    contenido = GeneradorExcel.generar(
        titulo="Reporte Ventas",
        subtitulo="Periodo Mensual",
        columnas=cols,
        filas=filas,
        totales=totales,
        metadatos=meta,
    )

    assert isinstance(contenido, bytes)
    assert len(contenido) > 500

    # Comprobar legibilidad
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    assert "Reporte Ejecutivo" in wb.sheetnames
    ws = wb["Reporte Ejecutivo"]
    assert ws.cell(row=5, column=1).value == "CODIGO"
    assert ws.cell(row=6, column=1).value == "V-001"


def test_generador_pdf_retorna_bytes_validos():
    """Valida que GeneradorPDF construya un flujo PDF institucional."""
    cols = ["CODIGO", "CLIENTE", "MONTO"]
    filas = [
        ["V-001", "Ana Valenzuela", 1250.50],
        ["V-002", "Carlos Mendez", 850.00],
    ]
    totales = {"MONTO": 2100.50}

    contenido = GeneradorPDF.generar(
        titulo="Reporte Ventas",
        subtitulo="Periodo Mensual",
        columnas=cols,
        filas=filas,
        totales=totales,
    )

    assert isinstance(contenido, bytes)
    assert len(contenido) > 1000
    assert contenido.startswith(b"%PDF-")


def test_generador_csv_retorna_bytes_validos_con_bom():
    """Valida que GeneradorCSV inyecte BOM UTF-8 y estructure filas delimitadas."""
    cols = ["CODIGO", "CLIENTE", "MONTO"]
    filas = [
        ["V-001", "Ana Valenzuela", 1250.50],
    ]

    contenido = GeneradorCSV.generar(
        titulo="Reporte Ventas",
        subtitulo="Periodo Mensual",
        columnas=cols,
        filas=filas,
    )

    assert isinstance(contenido, bytes)
    # Verificar prefijo BOM UTF-8 (\xef\xbb\xbf)
    assert contenido.startswith(b"\xef\xbb\xbf")
    texto = contenido.decode("utf-8")
    assert "V-001" in texto
    assert "Ana Valenzuela" in texto


# =========================================================================
# 2. Pruebas Unitarias del Parser de Comandos de Voz
# =========================================================================


def test_parser_voz_comando_ventas_excel():
    """Valida la identificacion de modulo ventas, formato excel y periodo este_mes."""
    comando = ParserComandosVoz.interpretar("Descargar reporte de ventas de este mes en excel")
    assert comando.modulo == ModuloReporteEnum.VENTAS
    assert comando.formato == FormatoReporteEnum.EXCEL
    assert comando.periodo == RangoTemporalEnum.ESTE_MES
    assert comando.intencion == IntencionVozEnum.EXPORTAR
    assert comando.accion_recomendada == "ejecutar_exportacion"


def test_parser_voz_comando_reservas_pdf_con_sucursal():
    """Valida identificacion de reservas, formato pdf y coincidencia de sucursal."""
    sucursales = [(10, "Equipetrol Boutique"), (20, "Calacoto Atelier")]
    comando = ParserComandosVoz.interpretar(
        "Consultar citas de fitting y reservas de equipetrol en pdf",
        sucursales_conocidas=sucursales,
    )
    assert comando.modulo == ModuloReporteEnum.RESERVAS
    assert comando.formato == FormatoReporteEnum.PDF
    assert comando.id_sucursal == 10
    assert comando.nombre_sucursal == "Equipetrol Boutique"
    assert comando.intencion == IntencionVozEnum.CONSULTAR


def test_parser_voz_comando_inventario_kardex_csv():
    """Valida sinonimos de kardex e inventario para hoy en csv."""
    comando = ParserComandosVoz.interpretar("Exportar stock de existencias de hoy en csv")
    assert comando.modulo == ModuloReporteEnum.INVENTARIO
    assert comando.formato == FormatoReporteEnum.CSV
    assert comando.periodo == RangoTemporalEnum.HOY


def test_parser_voz_comando_bitacora_consultar():
    """Valida modulo bitacora con intencion de consultar."""
    comando = ParserComandosVoz.interpretar("Ver logs de auditoria y accesos de esta semana")
    assert comando.modulo == ModuloReporteEnum.BITACORA
    assert comando.periodo == RangoTemporalEnum.ESTA_SEMANA
    assert comando.intencion == IntencionVozEnum.CONSULTAR
    assert comando.accion_recomendada == "actualizar_filtros"


def test_parser_voz_sucursal_consolidada_todas():
    """Valida que menciones de 'todas' o 'consolidado' anulen el filtro de sucursal."""
    sucursales = [(10, "Equipetrol"), (20, "Calacoto")]
    comando = ParserComandosVoz.interpretar(
        "Reporte global de ventas de todas las sucursales",
        sucursales_conocidas=sucursales,
    )
    assert comando.modulo == ModuloReporteEnum.VENTAS
    assert comando.id_sucursal is None


def test_parser_voz_texto_ininteligible_lanza_error():
    """Valida que textos vacios o sin palabras clave lancen ComandoVozNoReconocidoError."""
    with pytest.raises(ComandoVozNoReconocidoError):
        ParserComandosVoz.interpretar("hola buenos dias como estas")


# =========================================================================
# 3. Pruebas Unitarias de Reglas RBAC en Servicio
# =========================================================================


def test_rbac_administrador_permite_acceso_total(usuario_admin: UsuarioORM):
    """El superadministrador puede acceder a cualquier modulo y sucursal."""
    # No debe lanzar excepcion
    ServicioReportesVoz.validar_permisos(usuario_admin, ModuloReporteEnum.VENTAS, None)
    ServicioReportesVoz.validar_permisos(usuario_admin, ModuloReporteEnum.BITACORA, None)
    ServicioReportesVoz.validar_permisos(usuario_admin, ModuloReporteEnum.INVENTARIO, 10)


def test_rbac_encargado_bloquea_bitacora(usuario_encargado: UsuarioORM):
    """El encargado de sucursal tiene prohibido acceder a la bitacora corporativa."""
    with pytest.raises(AccesoReporteDenegadoError) as exc_info:
        ServicioReportesVoz.validar_permisos(usuario_encargado, ModuloReporteEnum.BITACORA, 10)
    assert "bitacora" in str(exc_info.value.message).lower()


def test_rbac_encargado_bloquea_sucursal_ajena_o_global(usuario_encargado: UsuarioORM):
    """El encargado solo puede consultar su propia sucursal asignada."""
    # Bloqueo de consolidado global
    with pytest.raises(AccesoReporteDenegadoError):
        ServicioReportesVoz.validar_permisos(usuario_encargado, ModuloReporteEnum.VENTAS, None)

    # Bloqueo de otra sucursal
    with pytest.raises(AccesoReporteDenegadoError):
        ServicioReportesVoz.validar_permisos(usuario_encargado, ModuloReporteEnum.VENTAS, 99)

    # Acceso permitido a su propia sede
    ServicioReportesVoz.validar_permisos(usuario_encargado, ModuloReporteEnum.VENTAS, 10)


def test_rbac_cajero_bloqueo_total(usuario_cajero: UsuarioORM):
    """Cajeros y clientes tienen bloqueo absoluto sobre el modulo de reportes."""
    with pytest.raises(AccesoReporteDenegadoError):
        ServicioReportesVoz.validar_permisos(usuario_cajero, ModuloReporteEnum.VENTAS, 10)


# =========================================================================
# 4. Pruebas de Integracion de Endpoints HTTP (FastAPI TestClient)
# =========================================================================


def test_endpoint_sin_token_retorna_401():
    """Peticiones sin token Bearer retornan HTTP 401."""
    app.dependency_overrides.clear()
    client = TestClient(app)
    resp = client.post("/api/v1/admin/reportes/previsualizar", json={"modulo": "ventas"})
    assert resp.status_code == 401


def test_endpoint_cajero_retorna_403(usuario_cajero: UsuarioORM, mock_db: MagicMock):
    """Cajeros autenticados reciben HTTP 403 al invocar reportes."""
    from core.deps import get_current_user
    app.dependency_overrides[get_db] = lambda: mock_db
    app.dependency_overrides[get_current_user] = lambda: usuario_cajero
    client = TestClient(app)
    try:
        resp = client.post(
            "/api/v1/admin/reportes/previsualizar",
            json={"modulo": "ventas", "formato": "excel", "periodo": "este_mes"},
            headers={"Authorization": "Bearer fake_token"},
        )
        assert resp.status_code == 403
    finally:
        app.dependency_overrides.clear()


def test_endpoint_interpretar_voz_exito(usuario_admin: UsuarioORM, mock_db: MagicMock):
    """Endpoint de interpretacion procesa el texto dictado y retorna comando estructurado."""
    from core.deps import get_current_user
    app.dependency_overrides[get_db] = lambda: mock_db
    app.dependency_overrides[get_current_user] = lambda: usuario_admin
    client = TestClient(app)
    try:
        resp = client.post(
            "/api/v1/admin/reportes/interpretar-voz",
            json={"texto_dictado": "descargar reporte de ventas en pdf"},
            headers={"Authorization": "Bearer fake_token"},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["modulo"] == "ventas"
        assert data["formato"] == "pdf"
        assert data["intencion"] == "exportar"
    finally:
        app.dependency_overrides.clear()


def test_endpoint_previsualizar_exito(usuario_admin: UsuarioORM, mock_db: MagicMock):
    """Endpoint de previsualizacion retorna conteo y sugerencia de nombre de archivo."""
    from core.deps import get_current_user
    app.dependency_overrides[get_db] = lambda: mock_db
    app.dependency_overrides[get_current_user] = lambda: usuario_admin
    client = TestClient(app)
    try:
        resp = client.post(
            "/api/v1/admin/reportes/previsualizar",
            json={"modulo": "ventas", "formato": "excel", "periodo": "este_mes"},
            headers={"Authorization": "Bearer fake_token"},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["modulo"] == "ventas"
        assert "total_registros" in data
        assert data["nombre_archivo_sugerido"].endswith(".xlsx")
    finally:
        app.dependency_overrides.clear()


def test_endpoint_exportar_streaming_post(usuario_admin: UsuarioORM, mock_db: MagicMock):
    """Endpoint POST /exportar transmite el streaming binario y cabecera Content-Disposition."""
    from core.deps import get_current_user
    app.dependency_overrides[get_db] = lambda: mock_db
    app.dependency_overrides[get_current_user] = lambda: usuario_admin
    client = TestClient(app)
    try:
        resp = client.post(
            "/api/v1/admin/reportes/exportar",
            json={"modulo": "ventas", "formato": "csv", "periodo": "este_mes"},
            headers={"Authorization": "Bearer fake_token"},
        )
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        assert "attachment; filename=" in resp.headers["content-disposition"]
        assert len(resp.content) > 0
    finally:
        app.dependency_overrides.clear()


def test_endpoint_exportar_streaming_get(usuario_admin: UsuarioORM, mock_db: MagicMock):
    """Endpoint GET /exportar permite la descarga directa por parametros de consulta."""
    from core.deps import get_current_user
    app.dependency_overrides[get_db] = lambda: mock_db
    app.dependency_overrides[get_current_user] = lambda: usuario_admin
    client = TestClient(app)
    try:
        resp = client.get(
            "/api/v1/admin/reportes/exportar?modulo=reservas&formato=pdf&periodo=este_mes",
            headers={"Authorization": "Bearer fake_token"},
        )
        assert resp.status_code == 200
        assert "application/pdf" in resp.headers["content-type"]
        assert resp.content.startswith(b"%PDF-")
    finally:
        app.dependency_overrides.clear()


def test_exportar_reporte_registra_evento_en_bitacora(usuario_admin: UsuarioORM, mock_db: MagicMock):
    """Valida que la exportacion registre un evento EXPORTAR_REPORTE en la bitacora de auditoria."""
    from unittest.mock import patch

    filtros = ReporteFiltrosIn(
        modulo=ModuloReporteEnum.VENTAS,
        formato=FormatoReporteEnum.EXCEL,
        periodo=RangoTemporalEnum.ESTE_MES,
    )

    with patch("modules.seguridad.cu30_bitacora.servicio.ServicioBitacoraAuditoria.registrar_evento_seguro") as mock_bitacora:
        contenido, nombre, media_type = ServicioReportesVoz.generar_archivo_reporte(
            db=mock_db,
            filtros=filtros,
            usuario_actual=usuario_admin,
        )

        assert isinstance(contenido, bytes)
        assert len(contenido) > 0
        assert mock_bitacora.called
        args, kwargs = mock_bitacora.call_args
        assert kwargs.get("accion") == "EXPORTAR_REPORTE"
        assert kwargs.get("tabla_modulo") == "reportes"
        payload = kwargs.get("payload_nuevo")
        assert payload is not None
        assert payload["modulo"] == "ventas"
        assert payload["formato"] == "excel"



